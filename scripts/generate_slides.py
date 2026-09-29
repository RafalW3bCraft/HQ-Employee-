#!/usr/bin/env python3
"""
Generate HQ-Employee presentation slides as both PPTX and PDF
derived strictly from submission/assemblyai/06_SLIDE_OUTLINE.md.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.colors import HexColor
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

OUTPUT_DIR = "/home/watcher/Desktop/employee/submission/assemblyai"
PPTX_PATH = os.path.join(OUTPUT_DIR, "HQ_Employee_Presentation.pptx")
PDF_PATH = os.path.join(OUTPUT_DIR, "HQ_Employee_Presentation.pdf")

SLIDES_DATA = [
    {
        "slide_num": 1,
        "title": "HQ-Employee",
        "subtitle": "The Governed AI Business Employee",
        "tagline": "Real-time 24 kHz speech discovery and client coordination bounded by deterministic commercial policy.",
        "bullets": [
            "Built on AssemblyAI's managed Voice Agent API (wss://agents.assemblyai.com/v1/ws)",
            "Native 24 kHz full-duplex speech with immediate barge-in interruption",
            "Deterministic tri-state Policy Engine (ALLOW / REQUIRE_APPROVAL / BLOCK)",
            "Zero secret leakage via single-use ephemeral HMAC voice tickets"
        ],
        "footer": "AssemblyAI Voice Agent Hackathon 2026 • Submission Deck"
    },
    {
        "slide_num": 2,
        "title": "The Enterprise Voice AI Challenge",
        "subtitle": "Why Traditional Conversational Bots Fail in Production",
        "tagline": "Unconstrained reasoning models cannot be trusted with commercial and legal authority.",
        "bullets": [
            "Hallucinated Pricing: LLMs invent discounts and project estimates without business boundaries",
            "Lack of Authority: No mathematical distinction between browsing info and signing contracts",
            "Turn-Taking Protocol Desync: Naive tool execution causes speech stutter, audio overlap, and desync",
            "Credential Leakage: Exposing raw API keys to browser clients violates enterprise security"
        ],
        "footer": "HQ-Employee • Enterprise Challenge"
    },
    {
        "slide_num": 3,
        "title": "Governed Architecture & Zero-Trust Security",
        "subtitle": "Single-Origin Delivery with Ephemeral Ticket Gating",
        "tagline": "The reasoning model is treated as an untrusted worker governed by deterministic gateway logic.",
        "bullets": [
            "Browser Console: AudioWorklet captures 24 kHz PCM16 mono audio with zero external dependencies",
            "Single-Use HMAC Tickets: Client requests 60s ticket (GET /api/voice/ticket); raw API keys stay server-side",
            "Strict Gateway Controls: Reused tickets (4003), expired tickets, invalid origins, and duplicate IPs (4029) rejected",
            "Append-Only Audit Log: Every action, decision, reason, timestamp, and actor ID recorded with credential sanitization"
        ],
        "footer": "HQ-Employee • Architecture & Security"
    },
    {
        "slide_num": 4,
        "title": "AssemblyAI Voice Agent API Integration",
        "subtitle": "Native 24 kHz Full-Duplex Speech & Strict Handshake",
        "tagline": "Direct WebSocket integration to wss://agents.assemblyai.com/v1/ws with bidirectional audio.",
        "bullets": [
            "Deterministic Handshake: session.update sent as first upstream message; session.ready confirms session_id",
            "Zero Audio Dropping: Microphone frames buffered before session.ready and drained immediately upon ready",
            "True Barge-In Interruption: voice.speech_started halts agent playback and flushes buffers instantly",
            "Low-Latency Neural Speech: Real-time turn detection and high-fidelity speech synthesis"
        ],
        "footer": "HQ-Employee • Voice Agent Engine"
    },
    {
        "slide_num": 5,
        "title": "Docs-Conformant Tool Coordination (BLK-012)",
        "subtitle": "Official Turn-Taking Protocol for Client-Side Tools",
        "tagline": "Tool execution synchronized to the exact AssemblyAI event specification.",
        "bullets": [
            "tool.call Arrives: Execution runs; result is buffered during active agent speech (reply.started)",
            "reply.done Arrives: Coordinator drains accumulated results and sends tool.result to upstream socket",
            "Interruption Protection: If speech is interrupted before reply.done, pending results are cleanly discarded",
            "13 Registered Business Tools: Strict JSON schemas governing discovery, budget, timeline, and calendar"
        ],
        "footer": "HQ-Employee • Tool Coordination Protocol"
    },
    {
        "slide_num": 6,
        "title": "Deterministic Policy Engine",
        "subtitle": "Fail-Closed Tri-State Commercial Boundaries",
        "tagline": "The AI reasoning engine never decides its own business or legal authority.",
        "bullets": [
            "ALLOW: Permitted routine operations (schedule_meeting, record_budget, record_timeline)",
            "REQUIRE_APPROVAL: Custom discount requests (e.g. 15% discount) escalate to human Commercial Director",
            "BLOCK: Legal agreements (sign_contract), wire transfers, and credential access strictly blocked",
            "Immutable Audit Trail: 100% of policy evaluations logged with decision provenance and timestamp"
        ],
        "footer": "HQ-Employee • Deterministic Governance"
    },
    {
        "slide_num": 7,
        "title": "Empirical Production Verification",
        "subtitle": "10 / 10 Live Integration Checks Passed",
        "tagline": "Grounded in real empirical outputs from docs/PRE_SUBMISSION_VERIFICATION.md.",
        "bullets": [
            "Health & Gateways: GET /health/live and /health/ready return 200 OK; single-use ticket lifecycle verified",
            "Audio & Protocol: 24 kHz WAV streaming confirmed with user transcript, agent speech, and reply.done",
            "Governance & Lifecycle: ALLOW/REQUIRE_APPROVAL/BLOCK verified; clean socket teardown on disconnect",
            "Test Suite & Security: 181 / 181 automated tests passing across 24 test suites; zero secret leaks"
        ],
        "footer": "HQ-Employee • Empirical Verification"
    },
    {
        "slide_num": 8,
        "title": "Deployment & Honest Disclosures",
        "subtitle": "Single-Origin Production Reality",
        "tagline": "Full transparency on runtime architecture, persistence, and external modalities.",
        "bullets": [
            "Google Cloud Run: Single-instance container (min=max=1, session-affinity, no-cpu-throttling)",
            "Persistence Model: Authoritative state in Node.js process memory; resets on container reboot (PERSISTENCE_STATUS.md)",
            "Database Schemas: Structured PostgreSQL migrations (001–006) for non-blocking write-through",
            "External Modalities: Google Calendar with simulated fallback; telephony SIP dispatch simulated; Android UI prototype"
        ],
        "footer": "HQ-Employee • Production Status & Disclosures"
    }
]


def create_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]  # blank layout

    # Colors
    c_bg = RGBColor(11, 15, 25)       # #0B0F19
    c_card = RGBColor(20, 27, 45)     # #141B2D
    c_cyan = RGBColor(56, 189, 248)   # #38BDF8
    c_indigo = RGBColor(129, 140, 248) # #818CF8
    c_white = RGBColor(241, 245, 249) # #F1F5F9
    c_muted = RGBColor(148, 163, 184) # #94A3B8
    c_accent_border = RGBColor(30, 41, 59)

    for data in SLIDES_DATA:
        slide = prs.slides.add_slide(blank_layout)

        # Background fill
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = c_bg
        bg_shape.line.fill.background()

        # Header Card Banner
        banner = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.6), Inches(11.733), Inches(1.8))
        banner.fill.solid()
        banner.fill.fore_color.rgb = c_card
        banner.line.color.rgb = c_indigo
        banner.line.width = Pt(1.5)

        # Slide Number Badge
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(11.2), Inches(0.8), Inches(1.1), Inches(0.45))
        badge.fill.solid()
        badge.fill.fore_color.rgb = c_indigo
        badge.line.fill.background()
        tf_b = badge.text_frame
        tf_b.text = f"0{data['slide_num']} / 08"
        p_b = tf_b.paragraphs[0]
        p_b.font.size = Pt(13)
        p_b.font.bold = True
        p_b.font.color.rgb = c_white
        p_b.alignment = PP_ALIGN.CENTER

        # Header Title
        tf_h = banner.text_frame
        tf_h.word_wrap = True
        p_title = tf_h.paragraphs[0]
        p_title.text = data["title"]
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = c_white

        p_sub = tf_h.add_paragraph()
        p_sub.text = data["subtitle"]
        p_sub.font.size = Pt(16)
        p_sub.font.color.rgb = c_cyan

        p_tag = tf_h.add_paragraph()
        p_tag.text = data["tagline"]
        p_tag.font.size = Pt(13)
        p_tag.font.italic = True
        p_tag.font.color.rgb = c_muted

        # Main Content Card
        content_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.7), Inches(11.733), Inches(4.0))
        content_box.fill.solid()
        content_box.fill.fore_color.rgb = c_card
        content_box.line.color.rgb = c_accent_border
        content_box.line.width = Pt(1)

        tf_c = content_box.text_frame
        tf_c.word_wrap = True

        for i, bullet in enumerate(data["bullets"]):
            p = tf_c.paragraphs[0] if i == 0 else tf_c.add_paragraph()
            p.text = f"  ✦   {bullet}"
            p.font.size = Pt(16)
            p.font.color.rgb = c_white
            p.space_after = Pt(18)

        # Footer
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.85), Inches(11.733), Inches(0.4))
        tf_f = footer_box.text_frame
        p_f = tf_f.paragraphs[0]
        p_f.text = data["footer"]
        p_f.font.size = Pt(11)
        p_f.font.color.rgb = c_muted

    prs.save(PPTX_PATH)
    print(f"Saved PPTX to: {PPTX_PATH}")


def create_pdf():
    # 16:9 Landscape page (792 x 445.5 pt ~ 11 x 6.1875 in)
    page_w = 792
    page_h = 445.5

    c = canvas.Canvas(PDF_PATH, pagesize=(page_w, page_h))

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'SlideTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=HexColor('#FFFFFF')
    )
    subtitle_style = ParagraphStyle(
        'SlideSub',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=HexColor('#38BDF8')
    )
    tagline_style = ParagraphStyle(
        'SlideTag',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=10,
        leading=14,
        textColor=HexColor('#94A3B8')
    )
    bullet_style = ParagraphStyle(
        'SlideBullet',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=11,
        leading=16,
        textColor=HexColor('#F1F5F9')
    )

    for data in SLIDES_DATA:
        # Background
        c.setFillColor(HexColor('#0B0F19'))
        c.rect(0, 0, page_w, page_h, fill=1, stroke=0)

        # Top Banner Card
        c.setFillColor(HexColor('#141B2D'))
        c.setStrokeColor(HexColor('#818CF8'))
        c.setLineWidth(1.5)
        c.roundRect(40, page_h - 130, page_w - 80, 100, 8, fill=1, stroke=1)

        # Slide Number Badge
        c.setFillColor(HexColor('#818CF8'))
        c.roundRect(page_w - 110, page_h - 60, 55, 22, 4, fill=1, stroke=0)
        c.setFillColor(HexColor('#FFFFFF'))
        c.setFont("Helvetica-Bold", 10)
        c.drawCentredString(page_w - 82.5, page_h - 53, f"0{data['slide_num']} / 08")

        # Titles
        p_t = Paragraph(data["title"], title_style)
        p_t.wrapOn(c, page_w - 180, 30)
        p_t.drawOn(c, 55, page_h - 62)

        p_s = Paragraph(data["subtitle"], subtitle_style)
        p_s.wrapOn(c, page_w - 180, 20)
        p_s.drawOn(c, 55, page_h - 84)

        p_tag = Paragraph(data["tagline"], tagline_style)
        p_tag.wrapOn(c, page_w - 180, 20)
        p_tag.drawOn(c, 55, page_h - 105)

        # Content Card
        c.setFillColor(HexColor('#141B2D'))
        c.setStrokeColor(HexColor('#1E293B'))
        c.setLineWidth(1)
        c.roundRect(40, 45, page_w - 80, page_h - 190, 8, fill=1, stroke=1)

        y_offset = page_h - 170
        for bullet in data["bullets"]:
            bullet_text = f"<b><font color='#38BDF8'>✦</font></b>&nbsp;&nbsp;{bullet}"
            p_b = Paragraph(bullet_text, bullet_style)
            w, h = p_b.wrap(page_w - 120, 50)
            p_b.drawOn(c, 60, y_offset - h)
            y_offset -= (h + 16)

        # Footer
        c.setFillColor(HexColor('#64748B'))
        c.setFont("Helvetica", 8)
        c.drawString(45, 25, data["footer"])
        c.drawRightString(page_w - 45, 25, "HQ-Employee • Governed AI Operating System")

        c.showPage()

    c.save()
    print(f"Saved PDF to: {PDF_PATH}")


if __name__ == "__main__":
    create_pptx()
    create_pdf()
